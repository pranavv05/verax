"""The register cleaner on the shape of register Indian clients actually send."""

import io

import pandas as pd
from openpyxl import load_workbook

from app.services.register_cleaner_service import _is_numeric_col, _parse_date, process_bytes


def _messy_register() -> bytes:
    rows = [
        ["ABC TRADERS PURCHASE REGISTER", None, None, None, None],
        [None, None, None, None, None],
        ["Date", "Party", "Invoice", "Taxable", "GST"],
        ["", None, "No.", "Value", "Amount"],
        ["01/09/2026", "Shree Steel", "SS-1", "1,20,000.00", "21,600"],
        ["05-09-2026", "Om Paper", "OP/22", "\u20b9 45,500", "8,190"],
        ["CANCELLED", "Om Paper", "OP/23", "-", "-"],
        ["12.09.2026", "Ravi Ent", "R9", "9800", "1764"],
        ["TOTAL", None, None, "175300", "31554"],
    ]
    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, header=False, index=False)
    return buf.getvalue()


def test_dates_are_read_day_first():
    assert _parse_date("01/09/2026") == pd.Timestamp(2026, 9, 1)
    assert _parse_date("12.09.2026") == pd.Timestamp(2026, 9, 12)
    assert _parse_date("5 Sep 2026") == pd.Timestamp(2026, 9, 5)


def test_amount_columns_are_recognised_but_identifiers_are_not():
    for col in ["Taxable Value", "Invoice Value", "Total", "GST Amount", "Bill Amt", "Net Amount"]:
        assert _is_numeric_col(col), col
    for col in ["Invoice No", "Tax Invoice No", "GSTIN", "GST No.", "Party"]:
        assert not _is_numeric_col(col), col


def test_cleans_a_messy_register():
    cleaned = pd.read_excel(io.BytesIO(process_bytes(_messy_register())))
    data = cleaned[cleaned["Party"].notna()]

    assert list(data["Date"]) == [
        pd.Timestamp(2026, 9, 1),
        pd.Timestamp(2026, 9, 5),
        pd.Timestamp(2026, 9, 12),
    ]
    assert list(data["Taxable Value"]) == [120000.0, 45500.0, 9800.0]
    assert list(data["GST Amount"]) == [21600.0, 8190.0, 1764.0]


def test_grand_total_sums_every_amount_column():
    sheet = load_workbook(io.BytesIO(process_bytes(_messy_register()))).active
    headers = [cell.value for cell in sheet[1]]
    total_row = [cell.value for cell in sheet[sheet.max_row]]

    assert total_row[0] == "GRAND TOTAL"
    for header in ("Taxable Value", "GST Amount"):
        formula = total_row[headers.index(header)]
        assert isinstance(formula, str) and formula.startswith("="), header
